"""
NETRA AI — Backend FastAPI (versi gratis / ringan)

Menyediakan analisis netnografi tanpa model ML berat:
- Unggah CSV/Excel berisi data media sosial -> ringkasan, sentimen, emosi,
  kata kunci, topik, distribusi platform.
- Analisis teks langsung (JSON).

Jalankan lokal:
    pip install -r requirements.txt
    uvicorn main:app --reload

Dokumentasi interaktif otomatis tersedia di /docs (Swagger UI).
"""

from __future__ import annotations

import io
from typing import List, Optional

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

import analysis

app = FastAPI(
    title="NETRA AI API",
    description="Backend netnografi ringan: sentimen, emosi, topik, kata kunci.",
    version="1.0.0",
)

# CORS: izinkan dipanggil dari frontend (Netlify) & localhost.
# Untuk produksi sebaiknya ganti "*" dengan domain spesifik Anda.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Kolom yang dikenali (fleksibel terhadap penamaan umum)
TEXT_COLS = ["content", "text", "teks", "tweet", "komentar", "isi", "caption"]
PLATFORM_COLS = ["platform", "sumber", "source", "media"]
AUTHOR_COLS = ["author", "user", "username", "akun", "penulis"]
DATE_COLS = ["date", "tanggal", "waktu", "created_at", "time"]
LIKES_COLS = ["likes", "like", "suka", "favorites"]
SHARES_COLS = ["shares", "share", "retweets", "bagikan"]


def _pick(columns: List[str], candidates: List[str]) -> Optional[str]:
    low = {c.lower().strip(): c for c in columns}
    for cand in candidates:
        if cand in low:
            return low[cand]
    return None


# --------------------------------------------------------------------------- #
# Skema request/response
# --------------------------------------------------------------------------- #

class TextRequest(BaseModel):
    texts: List[str]


class TextResultItem(BaseModel):
    text: str
    sentiment: str
    score: float
    emotion: str


@app.get("/")
def root():
    return {
        "name": "NETRA AI API",
        "status": "ok",
        "endpoints": ["/health", "/api/analyze-text", "/api/analyze-file", "/docs"],
    }


@app.get("/health")
def health():
    return {"status": "healthy"}


@app.post("/api/analyze-text", response_model=List[TextResultItem])
def analyze_text(req: TextRequest):
    """Analisis daftar teks secara langsung (tanpa unggah file)."""
    if not req.texts:
        raise HTTPException(status_code=400, detail="Daftar teks kosong.")
    results = []
    for t in req.texts:
        label, score = analysis.analyze_sentiment(t)
        results.append(TextResultItem(
            text=t, sentiment=label, score=score, emotion=analysis.detect_emotion(t)
        ))
    return results


@app.post("/api/analyze-file")
async def analyze_file(file: UploadFile = File(...)):
    """
    Unggah CSV atau Excel berisi data media sosial.
    Minimal harus ada kolom teks (content/text/teks/komentar/...).
    Mengembalikan ringkasan analisis siap dipakai dashboard.

    Dibaca tanpa pandas (stdlib csv + openpyxl) agar instalasi ringan & bebas
    kompilasi di Python versi apa pun.
    """
    raw = await file.read()
    name = (file.filename or "").lower()

    try:
        if name.endswith((".xlsx", ".xls")):
            rows = _read_excel(raw)
        else:
            rows = _read_csv(raw)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file: {e}")

    if not rows:
        raise HTTPException(status_code=400, detail="File tidak berisi data.")

    cols = list(rows[0].keys())
    text_col = _pick(cols, TEXT_COLS)
    if not text_col:
        raise HTTPException(
            status_code=400,
            detail=f"Kolom teks tidak ditemukan. Gunakan salah satu: {TEXT_COLS}. Kolom Anda: {cols}",
        )

    platform_col = _pick(cols, PLATFORM_COLS)
    author_col = _pick(cols, AUTHOR_COLS)
    date_col = _pick(cols, DATE_COLS)
    likes_col = _pick(cols, LIKES_COLS)
    shares_col = _pick(cols, SHARES_COLS)

    texts = [str(r.get(text_col) or "") for r in rows]

    # Sentimen & emosi per baris
    sent_counts = {"positif": 0, "netral": 0, "negatif": 0}
    emotion_counts: dict = {}
    platform_counter: dict = {}
    authors: set = set()
    sample_posts = []

    for i, (t, row) in enumerate(zip(texts, rows)):
        label, score = analysis.analyze_sentiment(t)
        emo = analysis.detect_emotion(t)
        sent_counts[label] += 1
        emotion_counts[emo] = emotion_counts.get(emo, 0) + 1

        if platform_col:
            pf = str(row.get(platform_col) or "—")
            platform_counter[pf] = platform_counter.get(pf, 0) + 1
        if author_col:
            authors.add(str(row.get(author_col) or ""))

        if len(sample_posts) < 20:
            sample_posts.append({
                "platform": str(row.get(platform_col)) if platform_col else "—",
                "author": str(row.get(author_col)) if author_col else "anonim",
                "content": t[:280],
                "date": str(row.get(date_col)) if date_col else "",
                "likes": _to_int(row.get(likes_col)) if likes_col else 0,
                "shares": _to_int(row.get(shares_col)) if shares_col else 0,
                "sentiment": label,
                "emotion": emo,
            })

    total = len(texts)
    sentiment_pct = {
        k: round(v / total * 100, 1) for k, v in sent_counts.items()
    }

    platform_counts = [
        {"name": k, "value": v}
        for k, v in sorted(platform_counter.items(), key=lambda x: -x[1])
    ]

    keywords = analysis.top_keywords(texts, n=30)
    topics = analysis.simple_topics(keywords, k=5)
    emotions = [{"name": k, "value": v} for k, v in
                sorted(emotion_counts.items(), key=lambda x: -x[1])]

    return {
        "summary": {
            "totalPosts": total,
            "uniqueAuthors": len(authors) if author_col else None,
            "textColumn": text_col,
            "detectedColumns": {
                "platform": platform_col, "author": author_col,
                "date": date_col, "likes": likes_col, "shares": shares_col,
            },
        },
        "sentiment": sentiment_pct,
        "sentimentCounts": sent_counts,
        "emotions": emotions,
        "platformCounts": platform_counts,
        "keywords": keywords,
        "topics": topics,
        "samplePosts": sample_posts,
    }


def _read_csv(raw: bytes) -> List[dict]:
    """Baca CSV dari bytes; deteksi encoding & pemisah secara sederhana."""
    import csv

    text = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="ignore")

    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [dict(r) for r in reader]


def _read_excel(raw: bytes) -> List[dict]:
    """Baca sheet pertama Excel (.xlsx) via openpyxl menjadi list of dict."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h) if h is not None else f"col{i}"
                  for i, h in enumerate(next(rows_iter))]
    except StopIteration:
        return []
    out = []
    for r in rows_iter:
        if r is None or all(c is None for c in r):
            continue
        out.append({header[i]: r[i] if i < len(r) else None
                    for i in range(len(header))})
    return out



def _to_int(v) -> int:
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return 0
